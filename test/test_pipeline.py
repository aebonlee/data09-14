"""실행 흐름 — 가져오기 → 계산 → 판정 → DB 저장 → 보고서."""
import os
import shutil
import tempfile
import unittest

from helpers import GYRO_HEADER, VIB_HEADER, temp_settings, vib_rows, write_csv
from vibtool import db as dbm
from vibtool.pipeline import rejudge, run_inbox
from vibtool.report import collect, make_report

TH = {"version": "t1", "items": {"vib_*_rms": {"warn": 0.5, "fault": 1.0}}}
ZERO_GYRO = [(i * 0.001, 0, 0, 0) for i in range(4)]


class Pipeline(unittest.TestCase):
    def setUp(self):
        self.s, self.dir = temp_settings(thresholds=TH)
        self.conn = dbm.connect(self.s["db_path"])
        inbox = self.s["inbox_dir"]
        write_csv(inbox, "EQ-1_20260901-090000_vib.csv", VIB_HEADER, vib_rows([1, -1, 1, -1]))
        write_csv(inbox, "EQ-1_20260901-090000_gyro.csv", GYRO_HEADER, ZERO_GYRO)
        write_csv(inbox, "EQ-1_20260902-090000_vib.csv", VIB_HEADER, vib_rows([0.3, -0.3, 0.3, -0.3]))
        write_csv(inbox, "notes.csv", ("a",), [(1,)])

    def tearDown(self):
        self.conn.close()
        shutil.rmtree(self.dir, ignore_errors=True)

    def grade(self, day):
        return self.conn.execute(
            "SELECT m.status, j.grade, j.thresholds_version, j.algorithm_version FROM measurement m "
            "LEFT JOIN judgement j USING (measurement_id) WHERE measured_at LIKE ?", (day + "%",)).fetchone()

    def test_run_judge_store(self):
        r = run_inbox(self.conn, self.s)
        self.assertEqual((r["files_seen"], r["files_registered"], r["files_skipped"], r["files_failed"],
                          r["measurements_done"]), (4, 3, 0, 1, 1))
        g = self.grade("2026-09-01")
        self.assertEqual((g["status"], g["grade"], g["thresholds_version"]), ("완료", "고장", "t1"))
        self.assertTrue(g["algorithm_version"].startswith("algorithms.placeholder@"))
        # 결과 값: X rms = 1, Y rms = 0 (vib 9 + gyro 9 항목)
        v = dict(self.conn.execute("SELECT item, value FROM result").fetchall())
        self.assertEqual(len(v), 18)
        self.assertAlmostEqual(v["vib_X_rms"], 1.0)
        self.assertAlmostEqual(v["vib_Y_rms"], 0.0)
        self.assertEqual(self.grade("2026-09-02")["status"], "대기")
        # 미등록 장비는 자동 등록 + 경고
        self.assertTrue(any("미등록 장비" in m["내용"] for m in r["messages"]))
        log = self.conn.execute("SELECT * FROM run_log").fetchone()
        self.assertEqual((log["files_registered"], log["measurements_done"]), (3, 1))

    def test_rerun_skips_duplicates(self):
        run_inbox(self.conn, self.s)
        r = run_inbox(self.conn, self.s)
        self.assertEqual((r["files_registered"], r["files_skipped"], r["measurements_done"]), (0, 3, 0))
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM result").fetchone()[0], 18)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM source_file").fetchone()[0], 3)

    def test_late_pair_completes(self):
        run_inbox(self.conn, self.s)
        write_csv(self.s["inbox_dir"], "EQ-1_20260902-090000_gyro.csv", GYRO_HEADER, ZERO_GYRO)
        r = run_inbox(self.conn, self.s)
        self.assertEqual(r["measurements_done"], 1)
        # 내용이 똑같은 각속도 파일(모두 0)이 다른 측정 건에 있어도 등록하되 경고
        self.assertTrue(any("내용이 똑같은 파일" in m["내용"] for m in r["messages"]))
        g = self.grade("2026-09-02")
        self.assertEqual((g["status"], g["grade"]), ("완료", "정상"))   # rms 0.3 < 0.5

    def test_conflicting_file_not_overwritten(self):
        run_inbox(self.conn, self.s)
        other = tempfile.mkdtemp(dir=self.dir)
        write_csv(other, "EQ-1_20260901-090000_vib.csv", VIB_HEADER, vib_rows([5, -5, 5, -5]))
        r = run_inbox(self.conn, self.s, inbox=other)
        self.assertEqual(r["files_failed"], 1)
        self.assertIn("덮어쓰지 않았습니다", r["messages"][0]["내용"])
        self.assertEqual(self.grade("2026-09-01")["grade"], "고장")
        v = self.conn.execute("SELECT value FROM result WHERE item='vib_X_rms'").fetchone()[0]
        self.assertAlmostEqual(v, 1.0)

    def test_rejudge_with_new_threshold(self):
        run_inbox(self.conn, self.s)
        self.s["thresholds"] = {"version": "t2", "items": {"vib_*_rms": {"warn": 2.0, "fault": 3.0}}}
        self.assertEqual(rejudge(self.conn, self.s), 1)
        g = self.grade("2026-09-01")
        self.assertEqual((g["grade"], g["thresholds_version"]), ("정상", "t2"))

    def test_swappable_algorithm(self):
        self.s["algorithm"] = {"module": "fake_algorithm", "function": "run"}
        self.s["thresholds"] = {"version": "t3", "items": {"custom_x0": {"warn": 0.5, "fault": 2.0}}}
        dbm.upsert_equipment(self.conn, "EQ-1", "시험 장비", {"rpm": "1800"})
        run_inbox(self.conn, self.s)
        v = dict(self.conn.execute("SELECT item, value FROM result").fetchall())
        self.assertEqual(v, {"custom_x0": 1.0, "custom_rpm": 1800.0})
        g = self.grade("2026-09-01")
        self.assertEqual((g["grade"], g["algorithm_version"]), ("주의", "fake_algorithm@fake-9"))

    def test_algorithm_error_then_retry(self):
        # 가짜 알고리즘은 rpm 을 숫자로 바꾸므로 「모름」이면 실패합니다
        self.s["algorithm"] = {"module": "fake_algorithm", "function": "run"}
        self.s["thresholds"] = {"version": "t3", "items": {}}
        dbm.upsert_equipment(self.conn, "EQ-1", None, {"rpm": "모름"})
        r = run_inbox(self.conn, self.s)
        self.assertEqual(r["measurements_done"], 0)
        self.assertEqual(self.grade("2026-09-01")["status"], "오류")
        self.assertTrue(any(m["내용"].startswith("계산 실패") for m in r["messages"]))
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM result").fetchone()[0], 0)
        # 파라미터를 고치고 다시 실행하면 오류 건을 다시 계산합니다
        dbm.upsert_equipment(self.conn, "EQ-1", None, {"rpm": "1800"})
        r = run_inbox(self.conn, self.s)
        self.assertEqual(r["measurements_done"], 1)
        g = self.grade("2026-09-01")
        self.assertEqual((g["status"], g["grade"]), ("완료", "판정 불가"))

    def test_report(self):
        run_inbox(self.conn, self.s)
        d = collect(self.conn, self.s, "2026-09-01", "2026-09-01")
        self.assertEqual((len(d["done"]), len(d["other"]), d["counts"]["고장"]), (1, 0, 1))
        self.assertTrue(d["has_placeholder"])
        self.assertFalse(d["has_sample"])
        d2 = collect(self.conn, self.s, "2026-09-02", None)
        self.assertEqual((len(d2["done"]), len(d2["other"])), (0, 1))
        data, h, x = make_report(self.conn, self.s)
        with open(h, encoding="utf-8") as f:
            text = f.read()
        self.assertIn("임시 알고리즘", text)
        self.assertNotIn("<b>예시 데이터</b>", text)
        self.assertIn("vib_X_rms", text)
        self.assertNotIn("http", text.replace("http-equiv", ""))   # 외부 링크 없음(폐쇄망)
        if x:
            from openpyxl import load_workbook
            wb = load_workbook(x)
            self.assertEqual(wb.sheetnames, ["요약", "측정별판정", "결과값", "주의고장목록", "미완료"])
            self.assertEqual(wb["결과값"].max_row, 1 + 18)


if __name__ == "__main__":
    unittest.main()
