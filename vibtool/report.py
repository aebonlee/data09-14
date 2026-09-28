"""해석 결과보고서 — 기간·장비를 골라 HTML(파일 하나)과 엑셀로 만듭니다."""
import html
import os
from datetime import datetime, timedelta

from .judge import GRADES, NO_GRADE, fmt

PLACEHOLDER_PREFIX = "algorithms.placeholder"


def collect(conn, settings, date_from=None, date_to=None, equipment=None):
    where, args = ["1=1"], []
    if date_from:
        where.append("m.measured_at >= ?")
        args.append(date_from)
    if date_to:
        end = (datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1)).strftime("%Y-%m-%d")
        where.append("m.measured_at < ?")
        args.append(end)
    if equipment:
        where.append("m.equipment_id = ?")
        args.append(equipment)
    rows = conn.execute(
        "SELECT m.*, e.name AS equipment_name, j.grade, j.basis, j.thresholds_version, j.algorithm_version "
        "FROM measurement m JOIN equipment e ON e.equipment_id = m.equipment_id "
        "LEFT JOIN judgement j ON j.measurement_id = m.measurement_id "
        "WHERE " + " AND ".join(where) + " ORDER BY m.measured_at, m.equipment_id", args).fetchall()
    done, other = [], []
    for r in rows:
        item = dict(r)
        if r["status"] == "완료":
            item["results"] = [dict(x) for x in conn.execute(
                "SELECT item, value, unit FROM result WHERE measurement_id=? ORDER BY item",
                (r["measurement_id"],))]
            done.append(item)
        else:
            other.append(item)
    counts = {g: 0 for g in GRADES + [NO_GRADE]}
    for d in done:
        counts[d["grade"]] = counts.get(d["grade"], 0) + 1
    return {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "date_from": date_from or "처음", "date_to": date_to or "끝", "equipment": equipment or "전체",
        "done": done, "other": other, "counts": counts,
        "alerts": [d for d in done if d["grade"] in ("주의", "고장")],
        "has_sample": any(r["is_sample"] for r in rows),
        "has_placeholder": any((d["algorithm_version"] or "").startswith(PLACEHOLDER_PREFIX) for d in done),
        "algorithm_versions": sorted({d["algorithm_version"] for d in done}),
        "thresholds_versions": sorted({d["thresholds_version"] for d in done}),
        "phrases": settings.get("report_phrases", {}),
    }


CSS = """
:root{--ink:#1d2733;--muted:#56616e;--line:#d5dbe2;--bg:#ffffff;--soft:#f3f5f8;
--ok:#1f6f43;--warn:#8a5a00;--fault:#b3261e;--band:#fff4d6}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);font:16px/1.6 "Malgun Gothic","Apple SD Gothic Neo",sans-serif;
word-break:keep-all;overflow-wrap:break-word}
main{max-width:calc(1100px + 2*var(--pad,20px));margin:0 auto;padding:24px var(--pad,20px) 48px}
@media (min-width:640px){main{--pad:28px}}
h1{font-size:1.5rem;margin:0 0 4px}h2{font-size:1.15rem;margin:32px 0 10px;border-bottom:2px solid var(--line);padding-bottom:4px}
.meta{color:var(--muted);margin:0 0 16px}
.band{background:var(--band);border:1px solid #e8c766;border-radius:6px;padding:10px 14px;margin:0 0 12px}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(140px,100%),1fr));gap:10px}
.card{border:1px solid var(--line);border-radius:6px;padding:10px 12px;background:var(--soft)}
.card b{display:block;font-size:1.6rem}
.tbl{overflow-x:auto;border:1px solid var(--line);border-radius:6px}
table{border-collapse:collapse;width:100%;font-size:.93rem}
th,td{border-bottom:1px solid var(--line);padding:6px 8px;text-align:left;vertical-align:top}
th{background:var(--soft);white-space:nowrap}td.num{text-align:right;font-variant-numeric:tabular-nums}
.g-정상{color:var(--ok);font-weight:700}.g-주의{color:var(--warn);font-weight:700}.g-고장{color:var(--fault);font-weight:700}
details{border:1px solid var(--line);border-radius:6px;margin:8px 0;padding:6px 10px}
summary{cursor:pointer}
@media print{details{break-inside:avoid}details:not([open]) > *:not(summary){display:block}}
"""


def _e(v):
    return html.escape("" if v is None else str(v))


def _g(grade):
    return '<span class="g-%s">%s</span>' % (_e(grade), _e(grade))


def render_html(data):
    p = []
    p.append('<!doctype html><html lang="ko"><head><meta charset="utf-8">'
             '<meta name="viewport" content="width=device-width,initial-scale=1">'
             '<title>해석 결과보고서</title><style>%s</style></head><body><main>' % CSS)
    p.append("<h1>진동·각속도 해석 결과보고서</h1>")
    p.append('<p class="meta">기간 %s ~ %s · 장비 %s · 작성 %s</p>' % (
        _e(data["date_from"]), _e(data["date_to"]), _e(data["equipment"]), _e(data["generated_at"])))
    if data["has_sample"]:
        p.append('<p class="band"><b>예시 데이터</b> — 이 보고서에는 시연용 가상 측정 파일로 만든 결과가 들어 있습니다. 실제 장비 상태가 아닙니다.</p>')
    if data["has_placeholder"]:
        p.append('<p class="band"><b>임시 알고리즘</b> — 결과 값은 기개발 알고리즘이 아니라 시험용 기본 통계(RMS·최대값·첨도)이고, 판정 기준도 가상의 값입니다. 흐름 확인용이며 판정 근거로 쓰지 않습니다.</p>')
    p.append('<p class="meta">알고리즘 버전: %s<br>판정 기준 버전: %s</p>' % (
        _e(", ".join(data["algorithm_versions"]) or "-"), _e(", ".join(data["thresholds_versions"]) or "-")))

    p.append("<h2>판정 요약</h2><div class=\"cards\">")
    for g, n in data["counts"].items():
        if g == NO_GRADE and n == 0:
            continue
        p.append('<div class="card">%s<b>%d건</b></div>' % (_g(g), n))
    p.append('<div class="card">미완료(대기·오류)<b>%d건</b></div></div>' % len(data["other"]))

    p.append("<h2>주의·고장 건</h2>")
    if data["alerts"]:
        p.append('<div class="tbl"><table><tr><th>측정 일시</th><th>장비</th><th>판정</th><th>근거</th><th>해석</th></tr>')
        for d in data["alerts"]:
            p.append("<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>" % (
                _e(d["measured_at"]), _e(d["equipment_id"]), _g(d["grade"]),
                _e(d["basis"]).replace("; ", "<br>"), _e(data["phrases"].get(d["grade"], ""))))
        p.append("</table></div>")
    else:
        p.append("<p>해당 기간에 주의·고장 건이 없습니다.</p>")

    p.append("<h2>측정 건별 판정</h2>")
    p.append('<div class="tbl"><table><tr><th>측정 일시</th><th>장비</th><th>장비명</th><th>판정</th><th>해석</th></tr>')
    for d in data["done"]:
        p.append("<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>" % (
            _e(d["measured_at"]), _e(d["equipment_id"]), _e(d["equipment_name"] or "-"), _g(d["grade"]),
            _e(data["phrases"].get(d["grade"], ""))))
    p.append("</table></div>")

    p.append("<h2>측정 건별 결과 값</h2>")
    for d in data["done"]:
        p.append("<details><summary>%s · %s · %s</summary>" % (
            _e(d["measured_at"]), _e(d["equipment_id"]), _g(d["grade"])))
        p.append('<div class="tbl"><table><tr><th>결과 항목</th><th>값</th><th>단위</th></tr>')
        for r in d["results"]:
            p.append('<tr><td>%s</td><td class="num">%s</td><td>%s</td></tr>' % (
                _e(r["item"]), _e(fmt(r["value"])), _e(r["unit"])))
        p.append("</table></div></details>")

    if data["other"]:
        p.append("<h2>미완료 측정 건</h2>")
        p.append('<div class="tbl"><table><tr><th>측정 일시</th><th>장비</th><th>상태</th><th>사유</th></tr>')
        for d in data["other"]:
            p.append("<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>" % (
                _e(d["measured_at"]), _e(d["equipment_id"]), _e(d["status"]), _e(d["status_note"])))
        p.append("</table></div>")
    p.append("</main></body></html>")
    return "\n".join(p)


def write_xlsx(data, path):
    """openpyxl 이 없으면 False 를 돌려줍니다(HTML 보고서는 그대로 만들어짐)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws.append(["진동·각속도 해석 결과보고서"])
    ws.append(["기간", "%s ~ %s" % (data["date_from"], data["date_to"])])
    ws.append(["장비", data["equipment"]])
    ws.append(["작성", data["generated_at"]])
    ws.append(["알고리즘 버전", ", ".join(data["algorithm_versions"])])
    ws.append(["판정 기준 버전", ", ".join(data["thresholds_versions"])])
    if data["has_sample"]:
        ws.append(["예시 데이터", "시연용 가상 측정 파일 결과가 들어 있습니다. 실제 장비 상태가 아닙니다."])
    if data["has_placeholder"]:
        ws.append(["임시 알고리즘", "시험용 기본 통계와 가상 기준입니다. 판정 근거로 쓰지 않습니다."])
    ws.append([])
    ws.append(["판정", "건수"])
    for g, n in data["counts"].items():
        ws.append([g, n])
    ws.append(["미완료(대기·오류)", len(data["other"])])
    ws["A1"].font = Font(bold=True, size=14)

    ws2 = wb.create_sheet("측정별판정")
    ws2.append(["측정 일시", "장비", "장비명", "판정", "근거", "해석", "판정 기준 버전", "알고리즘 버전"])
    for d in data["done"]:
        ws2.append([d["measured_at"], d["equipment_id"], d["equipment_name"], d["grade"], d["basis"],
                    data["phrases"].get(d["grade"], ""), d["thresholds_version"], d["algorithm_version"]])
    ws3 = wb.create_sheet("결과값")
    ws3.append(["측정 일시", "장비", "결과 항목", "값", "단위"])
    for d in data["done"]:
        for r in d["results"]:
            ws3.append([d["measured_at"], d["equipment_id"], r["item"], r["value"], r["unit"]])
    ws4 = wb.create_sheet("주의고장목록")
    ws4.append(["측정 일시", "장비", "판정", "근거"])
    for d in data["alerts"]:
        ws4.append([d["measured_at"], d["equipment_id"], d["grade"], d["basis"]])
    ws5 = wb.create_sheet("미완료")
    ws5.append(["측정 일시", "장비", "상태", "사유"])
    for d in data["other"]:
        ws5.append([d["measured_at"], d["equipment_id"], d["status"], d["status_note"]])
    for s in (ws2, ws3, ws4, ws5):
        for c in s[1]:
            c.font = Font(bold=True)
        for col, width in zip("ABCDEFGH", (20, 12, 16, 12, 50, 40, 20, 30)):
            s.column_dimensions[col].width = width
    wb.save(path)
    return True


def make_report(conn, settings, date_from=None, date_to=None, equipment=None, out_dir=None):
    data = collect(conn, settings, date_from, date_to, equipment)
    out_dir = out_dir or settings["report_dir"]
    os.makedirs(out_dir, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    base = "해석결과보고서_%s%s" % (stamp, "_예시데이터" if data["has_sample"] else "")
    html_path = os.path.join(out_dir, base + ".html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(render_html(data))
    xlsx_path = os.path.join(out_dir, base + ".xlsx")
    if not write_xlsx(data, xlsx_path):
        xlsx_path = None
    return data, html_path, xlsx_path
